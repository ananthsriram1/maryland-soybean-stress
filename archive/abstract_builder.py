from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new Document
doc = Document()

# Add Title
title = doc.add_paragraph("Quantifying Agricultural Resilience: A Geospatial Analysis of Drought, Soil Type, and the Mitigating Effect of Irrigation on Soybean Yield Disparity in Maryland")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.runs[0]
run.bold = True
run.font.size = Pt(12)
run.font.name = 'Times New Roman'

# Add space
doc.add_paragraph()

# Add Abstract Body
abstract_text = (
    "While the U.S. Midwest sets the global standard for soybean (Glycine max) productivity, "
    "Maryland consistently underperforms relative to the national benchmark. This state-level deficit, "
    "however, masks a critical intra-state disparity between the sandy Atlantic Coastal Plain (Eastern Shore) "
    "and the silt-loam Piedmont (North Central). This study quantifies agricultural resilience across these "
    "regions by investigating how drought exposure, soil type, and irrigation jointly shape soybean yield "
    "outcomes, with the central hypothesis that divergent resilience to hydrological stress is the primary "
    "driver of the yield gap. We employ a multi-factor geospatial framework integrating USDA NASS yield "
    "statistics, NOAA climate records, SSURGO soil data, and high-resolution Sentinel-2 imagery processed "
    "via Google Earth Engine. Focusing on the critical reproductive window (R4–R6), we analyze relationships "
    "between yield and satellite-derived vegetation and water indices (NDVI, NDWI) across five agricultural "
    "districts from 2017 to 2024. Results reveal a pedological paradox: districts on the Eastern Shore, "
    "characterized by coarse-textured soils with low water-holding capacity, exhibit greater yield stability "
    "than traditionally more favorable silt-loam districts, primarily due to more extensive irrigation "
    "infrastructure. These findings indicate that while soil texture historically conditioned drought "
    "vulnerability, active water management has begun to decouple production from climate variability in the "
    "Coastal Plain, leaving nominally stable rainfed zones increasingly exposed to seasonal flash droughts "
    "and reinforcing the need for targeted irrigation and drought-adaptation strategies in Maryland."
)

paragraph = doc.add_paragraph(abstract_text)
paragraph_format = paragraph.paragraph_format
paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
run = paragraph.runs[0]
run.font.size = Pt(11)
run.font.name = 'Times New Roman'

# Save the document
doc.save('Abstract_Submission.docx')